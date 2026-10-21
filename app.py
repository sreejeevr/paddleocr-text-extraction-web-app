import os

# =========================================================
# WINDOWS / PADDLE COMPATIBILITY
# =========================================================

os.environ["FLAGS_enable_pir_api"] = "0"

from flask import Flask, render_template, request, send_file
from werkzeug.utils import secure_filename

from paddleocr import PaddleOCR
import cv2

from docx import Document

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# =========================================================
# FLASK SETUP
# =========================================================

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"

os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)


# =========================================================
# OCR MODEL
# =========================================================

print()
print("=" * 55)
print("          AI OCR STUDIO")
print("=" * 55)
print("Loading OCR model...")
print("Please wait...")

ocr = PaddleOCR(
    lang="en",
    enable_mkldnn=False
)

print("OCR model ready.")
print("=" * 55)


# =========================================================
# IMAGE PREPARATION
# =========================================================

def prepare_image(image_path):

    image = cv2.imread(image_path)

    if image is None:

        raise ValueError(
            "Could not read the uploaded image."
        )

    height, width = image.shape[:2]

    # Only resize very large images
    max_width = 1600

    if width > max_width:

        scale = max_width / width

        new_width = int(width * scale)
        new_height = int(height * scale)

        image = cv2.resize(
            image,
            (new_width, new_height),
            interpolation=cv2.INTER_AREA
        )

    return image


# =========================================================
# SORT OCR RESULTS
# =========================================================

def sort_ocr_results(
    texts,
    scores,
    boxes
):

    data = []

    for text, score, box in zip(
        texts,
        scores,
        boxes
    ):

        try:

            x1 = float(box[0])
            y1 = float(box[1])
            x2 = float(box[2])
            y2 = float(box[3])

            center_x = (
                x1 + x2
            ) / 2

            center_y = (
                y1 + y2
            ) / 2

            data.append(
                {
                    "text": str(text),
                    "score": float(score),
                    "center_x": center_x,
                    "center_y": center_y
                }
            )

        except Exception:

            continue

    # Single or no result
    if len(data) <= 1:

        return data

    # Sort x positions
    x_positions = sorted(
        item["center_x"]
        for item in data
    )

    largest_gap = 0
    split_index = None

    # Find possible gap between columns
    for i in range(
        1,
        len(x_positions)
    ):

        gap = (
            x_positions[i]
            - x_positions[i - 1]
        )

        if gap > largest_gap:

            largest_gap = gap
            split_index = i

    # Single-column document
    if (
        split_index is None
        or largest_gap < 100
    ):

        data.sort(
            key=lambda item: (
                item["center_y"],
                item["center_x"]
            )
        )

        return data

    # Two-column document
    split_x = (
        x_positions[split_index - 1]
        + x_positions[split_index]
    ) / 2

    left_column = []
    right_column = []

    for item in data:

        if item["center_x"] < split_x:

            left_column.append(
                item
            )

        else:

            right_column.append(
                item
            )

    # Top to bottom
    left_column.sort(
        key=lambda item: item["center_y"]
    )

    right_column.sort(
        key=lambda item: item["center_y"]
    )

    return (
        left_column
        + right_column
    )


# =========================================================
# RUN OCR
# =========================================================

def run_ocr(image_path):

    print()
    print("-" * 55)
    print(
        f"OCR processing: {image_path}"
    )

    # Read image
    image = prepare_image(
        image_path
    )

    height, width = image.shape[:2]

    print(
        f"Image size: {width} x {height}"
    )

    # Run PaddleOCR
    results = ocr.predict(
        image
    )

    all_text = []
    all_scores = []
    all_boxes = []

    # Extract OCR results
    for result in results:

        try:

            texts = result["rec_texts"]
            scores = result["rec_scores"]
            boxes = result["rec_boxes"]

        except Exception:

            continue

        all_text.extend(
            texts
        )

        all_scores.extend(
            scores
        )

        for box in boxes:

            all_boxes.append(
                box
            )

    # No text found
    if not all_text:

        print(
            "No text detected."
        )

        print("-" * 55)

        return "", 0

    # Correct reading order
    ordered = sort_ocr_results(
        all_text,
        all_scores,
        all_boxes
    )

    final_text = []
    final_scores = []

    for item in ordered:

        text = item["text"].strip()

        if text:

            final_text.append(
                text
            )

            final_scores.append(
                item["score"]
            )

    # Join text
    text = "\n".join(
        final_text
    )

    # Calculate confidence
    if final_scores:

        confidence = (
            sum(final_scores)
            / len(final_scores)
        ) * 100

    else:

        confidence = 0

    print(
        f"Detected lines: {len(final_text)}"
    )

    print(
        f"Average confidence: "
        f"{confidence:.2f}%"
    )

    print("-" * 55)

    return text, confidence


# =========================================================
# HOME PAGE
# =========================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# =========================================================
# IMAGE / CAMERA OCR
# =========================================================

@app.route(
    "/ocr",
    methods=["POST"]
)
def perform_ocr():

    # Check image
    if "image" not in request.files:

        return render_template(
            "index.html",
            error="No image selected.",
            active_tab="ocr"
        )

    file = request.files[
        "image"
    ]

    # Empty filename
    if file.filename == "":

        return render_template(
            "index.html",
            error="Please select or capture an image.",
            active_tab="ocr"
        )

    # Secure filename
    filename = secure_filename(
        file.filename
    )

    # Allowed image formats
    allowed_extensions = (
        ".jpg",
        ".jpeg",
        ".png",
        ".bmp",
        ".webp"
    )

    if not filename.lower().endswith(
        allowed_extensions
    ):

        return render_template(
            "index.html",
            error=(
                "Supported formats: "
                "JPG, JPEG, PNG, BMP and WEBP."
            ),
            active_tab="ocr"
        )

    # Save uploaded/captured image
    image_path = os.path.join(
        UPLOAD_FOLDER,
        filename
    )

    file.save(
        image_path
    )

    try:

        # Run OCR
        text, confidence = run_ocr(
            image_path
        )

        # No text
        if not text:

            return render_template(
                "index.html",
                error=(
                    "No text detected. "
                    "Try a clearer image."
                ),
                active_tab="ocr"
            )

        # Display result
        return render_template(
            "index.html",

            text=text,

            confidence=(
                f"{confidence:.2f}"
            ),

            filename=filename,

            active_tab="ocr"
        )

    except Exception as e:

        print()
        print("=" * 55)
        print("OCR ERROR")
        print("=" * 55)
        print(e)
        print("=" * 55)

        return render_template(
            "index.html",

            error=(
                f"OCR Error: {e}"
            ),

            active_tab="ocr"
        )


# =========================================================
# DOWNLOAD EXTRACTED TEXT AS TXT
# =========================================================

@app.route(
    "/download-txt",
    methods=["POST"]
)
def download_txt():

    text = request.form.get(
        "text",
        ""
    )

    if not text.strip():

        return "No text available.", 400

    txt_path = os.path.join(
        UPLOAD_FOLDER,
        "extracted_text.txt"
    )

    with open(
        txt_path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            text
        )

    return send_file(
        txt_path,

        as_attachment=True,

        download_name=(
            "extracted_text.txt"
        ),

        mimetype="text/plain"
    )


# =========================================================
# DOWNLOAD EXTRACTED TEXT AS WORD
# =========================================================

@app.route(
    "/download-word",
    methods=["POST"]
)
def download_word():

    text = request.form.get(
        "text",
        ""
    )

    if not text.strip():

        return "No text available.", 400

    word_path = os.path.join(
        UPLOAD_FOLDER,
        "extracted_text.docx"
    )

    document = Document()

    # Add heading
    document.add_heading(
        "Extracted OCR Text",
        level=1
    )

    # Add extracted text
    for line in text.splitlines():

        document.add_paragraph(
            line
        )

    # Save Word file
    document.save(
        word_path
    )

    return send_file(
        word_path,

        as_attachment=True,

        download_name=(
            "extracted_text.docx"
        ),

        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )


# =========================================================
# DOWNLOAD EXTRACTED TEXT AS PDF
# =========================================================

@app.route(
    "/download-pdf",
    methods=["POST"]
)
def download_pdf():

    text = request.form.get(
        "text",
        ""
    )

    if not text.strip():

        return "No text available.", 400

    pdf_path = os.path.join(
        UPLOAD_FOLDER,
        "extracted_text.pdf"
    )

    # Create PDF
    pdf = canvas.Canvas(
        pdf_path,
        pagesize=A4
    )

    page_width, page_height = A4

    # Margins
    left_margin = 50
    right_margin = 50

    x = left_margin
    y = page_height - 50

    # Font
    pdf.setFont(
        "Helvetica",
        11
    )

    # Maximum characters per line
    max_chars = 90

    # Add title
    pdf.setFont(
        "Helvetica-Bold",
        16
    )

    pdf.drawString(
        x,
        y,
        "Extracted OCR Text"
    )

    y -= 35

    pdf.setFont(
        "Helvetica",
        11
    )

    # Write text
    for original_line in text.splitlines():

        # Preserve empty lines
        if not original_line.strip():

            y -= 18

            continue

        words = original_line.split()

        current_line = ""

        for word in words:

            test_line = (
                current_line
                + " "
                + word
            ).strip()

            # Line too long
            if len(test_line) > max_chars:

                # New page
                if y < 50:

                    pdf.showPage()

                    pdf.setFont(
                        "Helvetica",
                        11
                    )

                    y = page_height - 50

                pdf.drawString(
                    x,
                    y,
                    current_line
                )

                y -= 18

                current_line = word

            else:

                current_line = (
                    test_line
                )

        # Print remaining text
        if current_line:

            if y < 50:

                pdf.showPage()

                pdf.setFont(
                    "Helvetica",
                    11
                )

                y = page_height - 50

            pdf.drawString(
                x,
                y,
                current_line
            )

            y -= 18

    pdf.save()

    return send_file(
        pdf_path,

        as_attachment=True,

        download_name=(
            "extracted_text.pdf"
        ),

        mimetype="application/pdf"
    )


# =========================================================
# START FLASK SERVER
# =========================================================

if __name__ == "__main__":

    print()
    print("=" * 55)
    print("        AI OCR STUDIO SERVER")
    print("=" * 55)
    print(
        "Open: http://127.0.0.1:5000"
    )
    print("=" * 55)
    print()

    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )